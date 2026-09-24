"""
Generate Project Proposal Document in Word Format (3-4 Pages)
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime

# Create document
doc = Document()

# Set document margins
sections = doc.sections
for section in sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Title Page
title = doc.add_heading('PROJECT PROPOSAL', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_heading('AI-Powered Hotel Management System', level=2)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()
doc.add_paragraph()

# Project Info
info_para = doc.add_paragraph()
info_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
info_para.add_run('Submitted By:\n').bold = True
info_para.add_run('Student Name\n')
info_para.add_run('Registration Number\n')
info_para.add_run('Department of Computer Science\n')
info_para.add_run(f'\n{datetime.now().strftime("%B %Y")}')

doc.add_page_break()

# 1. Introduction
doc.add_heading('1. INTRODUCTION', level=1)
doc.add_paragraph(
    'The hospitality industry is experiencing rapid digital transformation, with hotels '
    'seeking innovative solutions to enhance customer experience and operational efficiency. '
    'This project proposes an AI-Powered Hotel Management System that integrates artificial '
    'intelligence and machine learning to revolutionize hotel operations. The system combines '
    'intelligent room recommendation capabilities with an automated chatbot service, creating '
    'a seamless and personalized booking experience.'
)

doc.add_paragraph(
    'The system addresses two critical challenges: (1) helping customers find suitable rooms '
    'based on their preferences and budget using predictive analytics, and (2) providing '
    'instant, accurate responses to customer inquiries through natural language processing. '
    'By leveraging machine learning algorithms, the system analyzes user behavior, predicts '
    'preferences, and automates customer interactions effectively.'
)

# 2. Problem Statement
doc.add_heading('2. PROBLEM STATEMENT', level=1)

problems = [
    'Customers face overwhelming choices when selecting from numerous room types without '
    'personalized guidance, leading to decision fatigue and abandoned bookings.',
    
    'Traditional hotel systems show all rooms equally without considering individual customer '
    'preferences, budget constraints, travel purposes, or group size.',
    
    'Customer support is limited to business hours, and human agents cannot handle multiple '
    'queries simultaneously, causing delays and inconsistent information.',
    
    'Without intelligent matching between customer needs and room offerings, hotels experience '
    'lower booking conversion rates and reduced revenue.',
    
    'Hotels without AI capabilities face competitive disadvantages in providing personalized '
    'experiences and instant customer service.',
]

for problem in problems:
    doc.add_paragraph(problem, style='List Bullet')

# 3. Objectives
doc.add_heading('3. OBJECTIVES', level=1)
objectives = [
    'Develop an intelligent room recommendation system using machine learning to predict '
    'customer-room compatibility and booking likelihood.',
    
    'Create an AI-powered chatbot using natural language processing to understand queries '
    'and provide instant, accurate responses about hotel services.',
    
    'Design a comprehensive management platform integrating booking, user authentication, '
    'payment processing, and administrative controls.',
    
    'Implement real-time analytics to track user behavior, booking patterns, and model performance.',
    
    'Ensure system security, scalability, and user-friendly interface for customers and administrators.',
]

for obj in objectives:
    doc.add_paragraph(obj, style='List Number')


# 4. Scope
doc.add_heading('4. SCOPE OF THE PROJECT', level=1)

doc.add_heading('4.1 Included Features', level=2)
scope_in = [
    'AI recommendation engine with dual-model architecture for compatibility and booking prediction',
    'Intelligent chatbot with intent classification supporting 50+ intents',
    'User management with role-based access (customer and admin)',
    'Room management supporting multiple cities and 10 room types',
    'Booking system with payment integration and confirmation',
    'Analytics dashboard tracking performance metrics and revenue',
    'Real-time notifications and email confirmations',
]

for item in scope_in:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('4.2 Excluded from Scope', level=2)
scope_out = [
    'Mobile application development (web-only)',
    'Integration with third-party booking platforms',
    'Multi-language support (English only)',
    'Video/voice call support in chatbot',
]

for item in scope_out:
    doc.add_paragraph(item, style='List Bullet')

doc.add_page_break()

# 5. Proposed Methodology
doc.add_heading('5. PROPOSED METHODOLOGY', level=1)

doc.add_heading('5.1 AI Recommendation System', level=2)
doc.add_paragraph(
    'The recommendation system employs a dual-model machine learning architecture. '
    'Model 1 uses Random Forest Classification to predict room-user compatibility (High/Low match) '
    'based on user type, room type, budget, group size, and booking patterns. Model 2 uses '
    'Gradient Boosting Regression to predict booking probability (0-100%) considering factors '
    'like view time, previous bookings, and urgency.'
)


doc.add_paragraph(
    'The system processes 10 features: user type, room type, season, day type, booking advance, '
    'stay duration, group size, view time, previous bookings, and budget. Features undergo '
    'preprocessing (Label Encoding for categorical data, Standard Scaling for numerical data) '
    'before prediction. The overall recommendation score combines compatibility (60%) and '
    'booking likelihood (40%), prioritizing relevance while optimizing conversions.'
)

doc.add_heading('5.2 AI Chatbot System', level=2)
doc.add_paragraph(
    'The chatbot uses Natural Language Processing with a pipeline architecture combining '
    'TF-IDF (Term Frequency-Inverse Document Frequency) vectorization and Multinomial Naive '
    'Bayes classification. User queries undergo preprocessing: tokenization, lemmatization, '
    'and vectorization before intent prediction. The system recognizes 50+ intents covering '
    'room inquiries, pricing, booking procedures, amenities, and contact information.'
)

doc.add_paragraph(
    'Training involves creating comprehensive datasets with patterns and responses for each '
    'intent. The model achieves 95-99% accuracy through proper feature engineering and '
    'hyperparameter tuning. Confidence thresholding handles ambiguous queries, providing '
    'fallback responses when certainty is low.'
)

doc.add_heading('5.3 System Integration', level=2)
doc.add_paragraph(
    'The complete system integrates frontend user interfaces, backend APIs for business logic, '
    'AI model APIs for recommendations and chatbot, and database for persistent storage. '
    'RESTful APIs enable communication between components. Authentication uses JWT tokens, '
    'and payment processing integrates with secure payment gateways. Real-time analytics '
    'track user interactions, booking conversions, and model performance metrics.'
)


doc.add_page_break()

# 6. Expected Outcomes
doc.add_heading('6. EXPECTED OUTCOMES', level=1)
outcomes = [
    'Fully functional AI recommendation system achieving 85-99% accuracy in predicting room compatibility',
    'Intelligent chatbot with 95%+ intent classification accuracy providing instant customer support',
    'Increased booking conversion rates through personalized recommendations (expected 15-25% improvement)',
    'Reduced customer service workload by automating 70-80% of common inquiries',
    'Comprehensive analytics dashboard providing insights into user behavior and business performance',
    'Scalable architecture supporting multiple cities, room types, and concurrent users',
    'Enhanced customer satisfaction through personalized experiences and instant responses',
]

for outcome in outcomes:
    doc.add_paragraph(outcome, style='List Bullet')

# 7. Limitations
doc.add_heading('7. LIMITATIONS', level=1)
limitations = [
    'Initial training uses synthetic data; accuracy improves with real user data collection',
    'Chatbot limited to predefined intents; cannot handle completely novel query types',
    'Recommendation system requires minimum user input; cold start for completely new users',
    'System performance depends on data quality and regular model retraining',
    'Language support limited to English; localization requires additional NLP models',
]

for limitation in limitations:
    doc.add_paragraph(limitation, style='List Bullet')


# 8. Future Enhancements
doc.add_heading('8. FUTURE ENHANCEMENTS', level=1)
future = [
    'Implement deep learning models for improved prediction accuracy with larger datasets',
    'Add collaborative filtering to recommend rooms based on similar user preferences',
    'Develop mobile applications for iOS and Android platforms',
    'Integrate voice-based chatbot using speech recognition and synthesis',
    'Implement multi-language support for international customers',
    'Add sentiment analysis to gauge customer satisfaction from interactions',
    'Enable real-time pricing optimization based on demand and occupancy',
    'Integrate with third-party travel platforms and booking engines',
]

for item in future:
    doc.add_paragraph(item, style='List Bullet')

# 9. Conclusion
doc.add_heading('9. CONCLUSION', level=1)
doc.add_paragraph(
    'This project presents an innovative AI-Powered Hotel Management System that addresses '
    'critical challenges in the hospitality industry through intelligent automation. By combining '
    'machine learning-based room recommendations with natural language processing-powered chatbot '
    'services, the system creates a comprehensive solution for modern hotel management.'
)

doc.add_paragraph(
    'The dual-model recommendation engine ensures customers find rooms matching their preferences '
    'while optimizing booking conversions. The intelligent chatbot provides instant, accurate '
    'customer support 24/7, reducing operational costs and improving service quality. Together, '
    'these AI components create a competitive advantage for hotels while enhancing customer satisfaction.'
)

doc.add_paragraph(
    'The proposed system demonstrates the practical application of artificial intelligence in '
    'solving real-world business problems. Through careful design, robust methodology, and '
    'scalable architecture, this project delivers tangible value to both hotel operators and '
    'customers, paving the way for AI-driven transformation in the hospitality sector.'
)

# Save document
doc.save('Project_Proposal.docx')
print("✅ Project proposal generated successfully!")
print("📄 File saved as: Project_Proposal.docx")
